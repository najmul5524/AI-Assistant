import re
import datetime
from pathlib import Path
from typing import List, Any, Optional

import pymupdf

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import REPORTS_DIR, BOT_NAME, BASE_DIR

_font_path = BASE_DIR / "fonts" / "kalpurush.ttf"
FONTS_DIR = BASE_DIR / "fonts"
HAS_KALPURUSH = _font_path.exists()

def _has_bengali(text: str) -> bool:
    """Returns True if string contains Bengali Unicode characters."""
    return any('\u0980' <= c <= '\u09ff' for c in text)

def _markdown_to_html(title: str, text_content: str) -> str:
    """Converts structured markdown into high-definition HTML for PyMuPDF Story."""
    now_str = datetime.datetime.now().strftime("%d %B, %Y | %I:%M %p")
    
    html_body = []
    lines = text_content.split("\n")
    for line in lines:
        raw = line.strip()
        if not raw:
            continue
            
        formatted = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', raw)
        formatted = re.sub(r'\*(.*?)\*', r'<i>\1</i>', formatted)
        
        if raw.startswith("### "):
            h_text = formatted[4:].strip()
            html_body.append(f"<h3>{h_text}</h3>")
        elif raw.startswith("## "):
            h_text = formatted[3:].strip()
            html_body.append(f"<h2>{h_text}</h2>")
        elif raw.startswith("# "):
            h_text = formatted[2:].strip()
            html_body.append(f"<h2>{h_text}</h2>")
        elif raw.startswith("- ") or raw.startswith("* ") or raw.startswith("• "):
            b_text = formatted[2:].strip()
            html_body.append(f'<div class="bullet"><span class="dot">•</span> {b_text}</div>')
        elif raw.startswith("<"):
            html_body.append(raw)
        else:
            html_body.append(f"<p>{formatted}</p>")
            
    body_content = "\n".join(html_body)
    
    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
@font-face {{
    font-family: 'Kalpurush';
    src: url('kalpurush.ttf');
}}
body {{
    font-family: 'Kalpurush', sans-serif;
    font-size: 11pt;
    line-height: 1.65;
    color: #1e293b;
    margin: 0;
    padding: 0;
}}
.header-meta {{
    font-size: 8.5pt;
    color: #94a3b8;
    text-align: right;
    margin-bottom: 8px;
    letter-spacing: 0.5px;
}}
h1.doc-title {{
    color: #0f172a;
    font-size: 19pt;
    line-height: 1.35;
    margin: 0 0 4px 0;
    font-weight: bold;
}}
.date-bar {{
    font-size: 9pt;
    color: #64748b;
    margin-bottom: 10px;
}}
.divider {{
    height: 2px;
    background-color: #3b82f6;
    margin-bottom: 18px;
}}
h2 {{
    color: #0f172a;
    font-size: 14pt;
    margin-top: 18px;
    margin-bottom: 8px;
    font-weight: bold;
}}
h3 {{
    color: #1e3a8a;
    font-size: 12pt;
    margin-top: 14px;
    margin-bottom: 6px;
    font-weight: bold;
}}
p {{
    margin: 0 0 10px 0;
    text-align: justify;
}}
.bullet {{
    margin-left: 15px;
    margin-bottom: 6px;
    text-indent: -12px;
    padding-left: 12px;
}}
.dot {{
    color: #2563eb;
    font-size: 14pt;
    line-height: 0;
    vertical-align: middle;
    margin-right: 6px;
}}
.footer {{
    margin-top: 25px;
    padding-top: 10px;
    border-top: 1px solid #e2e8f0;
    font-size: 8.5pt;
    color: #94a3b8;
    font-style: italic;
}}
</style>
</head>
<body>
<div class="header-meta">{BOT_NAME} 24/7 AI • Confidential Executive Report</div>
<h1 class="doc-title">{title}</h1>
<div class="date-bar">জেনারেটেড বাই: {BOT_NAME} AI | সময়: {now_str}</div>
<div class="divider"></div>

{body_content}

<div class="footer">
গোপনীয় ও নির্ভরযোগ্য তথ্যসমৃদ্ধ — প্রস্তুত করেছে {BOT_NAME} পার্সোনাল এআই সহকারী।
</div>
</body>
</html>"""

def generate_pdf_report(title: str, text_content: str, filename_prefix: str = "report") -> Path:
    """
    Generates an executive-quality PDF report using PyMuPDF Story & Archive.
    Renders Bengali Unicode OpenType shaping (complex ligatures, conjuncts, and vowel reordering) 100% flawlessly.
    Returns the Path to the generated PDF.
    """
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    sanitized_prefix = re.sub(r'[^\w\-]', '_', filename_prefix)[:20]
    file_path = REPORTS_DIR / f"{sanitized_prefix}_{timestamp}.pdf"

    archive = pymupdf.Archive(str(FONTS_DIR))
    archive.add(str(REPORTS_DIR))
    html = _markdown_to_html(title, text_content)

    writer = pymupdf.DocumentWriter(str(file_path))
    story = pymupdf.Story(html=html, archive=archive)

    page_rect = pymupdf.Rect(0, 0, 595, 842) # Standard A4 page
    body_rect = pymupdf.Rect(40, 42, 595 - 40, 842 - 45) # Margins

    more = 1
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(body_rect)
        story.draw(device)
        writer.end_page()

    writer.close()
    return file_path

def generate_excel_report(title: str, headers: List[str], data_rows: List[List[Any]], filename_prefix: str = "report") -> Path:
    """
    Generates a beautifully formatted Excel report using OpenPyXL.
    Returns the Path to the generated .xlsx file.
    """
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    sanitized_prefix = re.sub(r'[^\w\-]', '_', filename_prefix)[:20]
    file_path = REPORTS_DIR / f"{sanitized_prefix}_{timestamp}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary Report"

    # Title Block
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = f"{title} - {BOT_NAME} Report"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32

    # Date Block
    ws.merge_cells("A2:E2")
    date_cell = ws["A2"]
    date_cell.value = f"Exported: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    date_cell.font = Font(name="Calibri", size=9, italic=True, color="64748B")
    date_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 18

    # Headers Block (Row 4)
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = str(header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 24

    # Data Rows
    row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    for r_idx, row_data in enumerate(data_rows, 5):
        is_even = (r_idx % 2 == 0)
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.value = val
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if is_even:
                cell.fill = row_alt_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r_idx].height = 20

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(str(file_path))
    return file_path
